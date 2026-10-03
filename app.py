import os
import io
import sqlite3
import hashlib
import base64
from datetime import datetime, timedelta, date
from functools import wraps

from flask import (
    Flask, render_template, request, redirect, url_for,
    session, jsonify, send_file, g
)
import bcrypt
from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from PIL import Image

from config import Config
from init_database import init_db
from face_engine.enrollment import enroll_face
from face_engine.recognition import recognize_face
from face_engine.liveness import check_liveness
from face_engine.utils import decode_base64_image

# ==========================================
# SYSTEM CORE SETUP (ABSOLUTE PATHS)
# ==========================================
BASE_DIR = os.path.dirname(os.path.abspath(__file__))

app = Flask(__name__)
app.secret_key = Config.SECRET_KEY
app.config["PERMANENT_SESSION_LIFETIME"] = timedelta(hours=Config.SESSION_TIMEOUT_HOURS)
app.config["MAX_CONTENT_LENGTH"] = 16 * 1024 * 1024

CLEAR_DATA_PASSWORD = "delete123"

# ==========================================
# TIME & DATABASE HELPERS
# ==========================================
def get_pkt_now():
    """Returns exact current datetime in Pakistan Standard Time (UTC+5)"""
    return datetime.utcnow() + timedelta(hours=5)

def get_db():
    if "db" not in g:
        db_path = Config.DATABASE
        if not os.path.isabs(db_path):
            db_path = os.path.join(BASE_DIR, db_path)
            
        g.db = sqlite3.connect(db_path)
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA journal_mode=WAL")
        g.db.execute("PRAGMA foreign_keys=ON")
        
        # Ensure absences table and photo columns exist automatically
        g.db.execute("""
            CREATE TABLE IF NOT EXISTS teacher_absences (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                teacher_id INTEGER,
                teacher_name TEXT,
                date TEXT,
                day_name TEXT,
                reason TEXT,
                timestamp DATETIME DEFAULT CURRENT_TIMESTAMP
            )
        """)
        try: g.db.execute("ALTER TABLE students ADD COLUMN photo_b64 TEXT")
        except Exception: pass
        try: g.db.execute("ALTER TABLE teachers ADD COLUMN photo_b64 TEXT")
        except Exception: pass

    return g.db

@app.teardown_appcontext
def close_db(exc):
    db = g.pop("db", None)
    if db is not None:
        db.close()

def get_setting(key, default=""):
    db = get_db()
    row = db.execute("SELECT value FROM settings WHERE key=?", (key,)).fetchone()
    return row["value"] if row else default

def log_audit(action, target_type, target_id, details, admin_user, ip):
    db = get_db()
    db.execute(
        "INSERT INTO audit_log (action,target_type,target_id,details,admin_user,ip_address) VALUES (?,?,?,?,?,?)",
        (action, target_type, target_id, details, admin_user, ip)
    )
    db.commit()

def save_face_photo(b64_str, folder, record_id):
    """Saves face photo dynamically to static uploads folder"""
    try:
        if not b64_str: return
        if "," in b64_str: b64_str = b64_str.split(",")[1]
        img_data = base64.b64decode(b64_str)
        img = Image.open(io.BytesIO(img_data)).convert("RGB")
        out_dir = os.path.join(BASE_DIR, "static", "uploads", folder)
        os.makedirs(out_dir, exist_ok=True)
        img.save(os.path.join(out_dir, f"{record_id}.jpg"), "JPEG", quality=90)
    except Exception as e:
        print(f"Error saving face photo file: {e}")

# ==========================================
# AUTHENTICATION
# ==========================================
def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if "user" not in session: return redirect(url_for("login"))
        if "login_time" in session:
            elapsed = datetime.now() - datetime.fromisoformat(session["login_time"])
            if elapsed > timedelta(hours=Config.SESSION_TIMEOUT_HOURS):
                session.clear()
                return redirect(url_for("login"))
        return f(*args, **kwargs)
    return decorated

def get_device_fingerprint(data):
    raw = f"{data.get('browser','')}{data.get('os','')}{data.get('screen','')}{data.get('canvas','')}"
    return hashlib.sha256(raw.encode()).hexdigest()

@app.route("/")
def index():
    if "user" in session: return redirect(url_for("dashboard"))
    return redirect(url_for("login"))

@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        data = request.get_json() if request.is_json else request.form
        username = (data.get("username") or "").strip()
        password = (data.get("password") or "").strip()
        ip = request.remote_addr
        db = get_db()
        
        lockout_row = db.execute("SELECT COUNT(*) as cnt FROM login_attempts WHERE ip_address=? AND success=0 AND timestamp > datetime('now', ?)", (ip, f"-{Config.LOCKOUT_MINUTES} minutes")).fetchone()
        if lockout_row["cnt"] >= Config.MAX_LOGIN_ATTEMPTS:
            return jsonify({"success": False, "error": f"Too many attempts. Locked for {Config.LOCKOUT_MINUTES} minutes."}), 429
            
        user = db.execute("SELECT * FROM users WHERE username=?", (username,)).fetchone()
        if user and bcrypt.checkpw(password.encode(), user["password_hash"].encode()):
            fp_data = data.get("fingerprint", {}) if isinstance(data, dict) else {}
            fp = get_device_fingerprint(fp_data)
            devices = db.execute("SELECT * FROM authorized_devices").fetchall()
            device_exists = any(d["device_fingerprint"] == fp for d in devices)
            
            if not device_exists:
                if len(devices) >= Config.MAX_DEVICES:
                    db.execute("INSERT INTO login_attempts (username,ip_address,success) VALUES (?,?,0)", (username, ip))
                    db.commit()
                    log_audit("BLOCKED_LOGIN", "security", None, f"Unknown device blocked for {username}", username, ip)
                    return jsonify({"success": False, "error": f"Device not authorized. Max {Config.MAX_DEVICES} devices reached."}), 403
                
                pkt_time_str = get_pkt_now().strftime('%b %d, %H:%M')
                ua = request.headers.get('User-Agent', '')
                if 'iPad' in ua or 'iPhone' in ua or 'Macintosh' in ua: label = f"📱 Apple iPad / iPhone ({pkt_time_str})"
                elif 'Android' in ua: label = f"📱 Android Phone ({pkt_time_str})"
                elif 'Windows' in ua: label = f"💻 Windows Computer ({pkt_time_str})"
                else: label = f"🌐 Web Device ({pkt_time_str})"
                db.execute("INSERT INTO authorized_devices (device_fingerprint, device_label) VALUES (?,?)", (fp, label))
                
            db.execute("INSERT INTO login_attempts (username,ip_address,success) VALUES (?,?,1)", (username, ip))
            db.execute("UPDATE authorized_devices SET last_seen=CURRENT_TIMESTAMP WHERE device_fingerprint=?", (fp,))
            db.commit()
            session.permanent = True
            session["user"] = username
            session["login_time"] = datetime.now().isoformat()
            return jsonify({"success": True, "redirect": url_for("dashboard")})
            
        db.execute("INSERT INTO login_attempts (username,ip_address,success) VALUES (?,?,0)", (username, ip))
        db.commit()
        return jsonify({"success": False, "error": "Invalid username or password."}), 401
    return render_template("login.html")

@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


# ==========================================
# DASHBOARD & ANALYTICS
# ==========================================
@app.route("/dashboard")
@login_required
def dashboard():
    db = get_db()
    pkt_now = get_pkt_now()
    today = pkt_now.date().isoformat()
    is_sunday = pkt_now.weekday() == 6
    holiday = db.execute("SELECT reason FROM holidays WHERE date=?", (today,)).fetchone()

    total_students = db.execute("SELECT COUNT(*) as c FROM students WHERE is_active=1").fetchone()["c"]
    total_teachers = db.execute("SELECT COUNT(*) as c FROM teachers WHERE is_active=1").fetchone()["c"]
    student_scans = db.execute("SELECT COUNT(*) as c FROM student_attendance WHERE date=?", (today,)).fetchone()["c"]
    teacher_scans = db.execute("SELECT COUNT(*) as c FROM teacher_attendance WHERE date=?", (today,)).fetchone()["c"]

    absent_students = []
    absent_teachers = []
    if not is_sunday and not holiday:
        # Check students consecutive absences
        all_students = db.execute("SELECT id, full_name, gr_number, class_name, father_name, phone FROM students WHERE is_active=1").fetchall()
        for s in all_students:
            consec = 0
            for offset in range(0, 5):
                d = (pkt_now.date() - timedelta(days=offset))
                if d.weekday() == 6: continue
                d_iso = d.isoformat()
                if db.execute("SELECT id FROM holidays WHERE date=?", (d_iso,)).fetchone(): continue
                if db.execute("SELECT id FROM student_attendance WHERE student_id=? AND date=?", (s["id"], d_iso)).fetchone(): break
                consec += 1
            if consec >= 2:
                absent_students.append({
                    "full_name": s["full_name"], "gr_number": s["gr_number"],
                    "class_name": s["class_name"], "father_name": s["father_name"],
                    "phone": s["phone"] or "", "consecutive_absent": consec
                })
        absent_students.sort(key=lambda x: x["consecutive_absent"], reverse=True)

        # Teachers absent today
        absent_teachers = db.execute("""
            SELECT t.id, t.full_name, t.employee_id, t.phone, t.designation
            FROM teachers t LEFT JOIN teacher_attendance ta ON ta.teacher_id = t.id AND ta.date = ?
            WHERE t.is_active = 1 AND ta.id IS NULL
        """, (today,)).fetchall()

    recent_scans = db.execute("""
        SELECT 'Student' as type, s.full_name, s.gr_number as id_num, s.class_name, sa.scan_time, sa.status
        FROM student_attendance sa JOIN students s ON s.id = sa.student_id WHERE sa.date = ?
        UNION ALL
        SELECT 'Teacher', t.full_name, t.employee_id, t.designation, ta.check_in, ta.status
        FROM teacher_attendance ta JOIN teachers t ON t.id = ta.teacher_id WHERE ta.date = ?
        ORDER BY scan_time DESC LIMIT 15
    """, (today, today)).fetchall()

    class_stats = db.execute("""
        SELECT s.class_name, COUNT(DISTINCT s.id) as total, COUNT(DISTINCT sa.student_id) as present
        FROM students s LEFT JOIN student_attendance sa ON sa.student_id = s.id AND sa.date = ?
        WHERE s.is_active = 1 GROUP BY s.class_name ORDER BY s.class_name
    """, (today,)).fetchall()

    week_trend = []
    for i in range(6, -1, -1):
        d = (pkt_now.date() - timedelta(days=i)).isoformat()
        cnt = db.execute("SELECT COUNT(*) as c FROM student_attendance WHERE date=?", (d,)).fetchone()["c"]
        week_trend.append({"date": d, "count": cnt})

    leaderboard_students = db.execute("SELECT s.full_name, s.class_name, COUNT(*) as days FROM student_attendance sa JOIN students s ON s.id = sa.student_id WHERE sa.date >= date('now','start of month') GROUP BY sa.student_id ORDER BY days DESC LIMIT 3").fetchall()
    leaderboard_teachers = db.execute("SELECT t.full_name, t.designation, COUNT(*) as days FROM teacher_attendance ta JOIN teachers t ON t.id = ta.teacher_id WHERE ta.date >= date('now','start of month') GROUP BY ta.teacher_id ORDER BY days DESC LIMIT 3").fetchall()

    return render_template("dashboard.html",
        total_students=total_students, total_teachers=total_teachers,
        student_scans=student_scans, teacher_scans=teacher_scans,
        present_now=student_scans + teacher_scans, is_sunday=is_sunday, holiday=holiday,
        absent_students=absent_students, absent_teachers=absent_teachers,
        recent_scans=recent_scans, class_stats=class_stats, week_trend=week_trend,
        leaderboard_students=leaderboard_students, leaderboard_teachers=leaderboard_teachers,
        user=session.get("user", ""), today=today
    )

@app.route("/api/live-stats")
@login_required
def live_stats():
    db = get_db()
    today = get_pkt_now().date().isoformat()
    ss = db.execute("SELECT COUNT(*) as c FROM student_attendance WHERE date=?", (today,)).fetchone()["c"]
    ts = db.execute("SELECT COUNT(*) as c FROM teacher_attendance WHERE date=?", (today,)).fetchone()["c"]
    return jsonify({"student_scans": ss, "teacher_scans": ts, "present": ss + ts})

@app.route("/api/analytics-filter")
@login_required
def analytics_filter():
    period = request.args.get("period", "day").lower()
    db = get_db()
    if period == "day": date_clause = "date = DATE('now')"
    elif period == "week": date_clause = "date >= DATE('now', '-7 days')"
    elif period == "month": date_clause = "date >= DATE('now', 'start of month')"
    elif period == "year": date_clause = "date >= DATE('now', 'start of year')"
    else: date_clause = "date = DATE('now')"
    
    s_count = db.execute(f"SELECT COUNT(*) as c FROM student_attendance WHERE {date_clause}").fetchone()["c"]
    t_count = db.execute(f"SELECT COUNT(*) as c FROM teacher_attendance WHERE {date_clause}").fetchone()["c"]
    return jsonify({"success": True, "period": period, "student_scans": s_count, "teacher_scans": t_count})


# ==========================================
# REGISTRATION
# ==========================================
@app.route("/register-student", methods=["GET", "POST"])
@login_required
def register_student():
    if request.method == "POST":
        data = request.get_json() or {}
        db = get_db()
        required = ["full_name", "father_name", "class_name", "section", "gr_number"]
        for k in required:
            if not (data.get(k) or "").strip(): return jsonify({"success": False, "error": f"Missing field: {k}"}), 400

        phone = (data.get("phone") or "").strip()
        images = data.get("images", [])
        if not images or len(images) < 3: return jsonify({"success": False, "error": "Capture all 3 angles."}), 400
        
        emb, _ = enroll_face(images)
        if emb is None: return jsonify({"success": False, "error": "No face detected in photos."}), 400
        
        try:
            primary_img = images[0]
            db.execute("INSERT INTO students (full_name, father_name, class_name, section, gr_number, phone, sibling_of, face_embedding, photo_b64) VALUES (?,?,?,?,?,?,?,?,?)",
                (data["full_name"].strip(), data["father_name"].strip(), data["class_name"], data["section"], data["gr_number"].strip(), phone, data.get("sibling_of"), emb, primary_img))
            db.commit()
            sid = db.execute("SELECT last_insert_rowid() as id").fetchone()["id"]
            save_face_photo(primary_img, "students", sid)
            log_audit("REGISTER", "student", sid, f"Registered student GR#{data['gr_number']} - {data['full_name']}", session["user"], request.remote_addr)
            return jsonify({"success": True, "id": sid})
        except sqlite3.IntegrityError: return jsonify({"success": False, "error": "GR Number already exists."}), 400
    return render_template("register_student.html", classes=Config.CLASSES, sections=Config.SECTIONS)

@app.route("/register-teacher", methods=["GET", "POST"])
@login_required
def register_teacher():
    if request.method == "POST":
        data = request.get_json() or {}
        db = get_db()
        for k in ["full_name", "father_name", "employee_id", "phone"]:
            if not (data.get(k) or "").strip(): return jsonify({"success": False, "error": f"Missing field: {k}"}), 400
            
        images = data.get("images", [])
        if not images or len(images) < 3: return jsonify({"success": False, "error": "Capture all 3 angles."}), 400
        
        emb, _ = enroll_face(images)
        if emb is None: return jsonify({"success": False, "error": "No face detected."}), 400
        
        try:
            primary_img = images[0]
            db.execute("INSERT INTO teachers (full_name, father_name, employee_id, subject, designation, department, joining_date, phone, email, face_embedding, photo_b64) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                (data["full_name"].strip(), data["father_name"].strip(), data["employee_id"].strip(), data.get("subject", ""), data.get("designation", ""), data.get("department", ""), data.get("joining_date", ""), data.get("phone", "").strip(), data.get("email", ""), emb, primary_img))
            db.commit()
            tid = db.execute("SELECT last_insert_rowid() as id").fetchone()["id"]
            save_face_photo(primary_img, "teachers", tid)
            log_audit("REGISTER", "teacher", tid, f"Registered teacher EMP#{data['employee_id']} - {data['full_name']}", session["user"], request.remote_addr)
            return jsonify({"success": True, "id": tid})
        except sqlite3.IntegrityError: return jsonify({"success": False, "error": "Employee ID already exists."}), 400
    return render_template("register_teacher.html")


# ==========================================
# MANAGE RECORDS (Students & Teachers)
# ==========================================
@app.route("/students")
@login_required
def students_page():
    db = get_db()
    students = db.execute("SELECT * FROM students WHERE is_active=1 ORDER BY class_name, section, full_name").fetchall()
    return render_template("students.html", students=students, count=len(students), classes=Config.CLASSES, sections=Config.SECTIONS)

@app.route("/student/<int:sid>/photo")
@login_required
def student_photo(sid):
    db = get_db()
    file_path = os.path.join(BASE_DIR, "static", "uploads", "students", f"{sid}.jpg")
    if os.path.exists(file_path): return send_file(file_path, mimetype="image/jpeg")
    try:
        row = db.execute("SELECT photo_b64 FROM students WHERE id=?", (sid,)).fetchone()
        if row and row["photo_b64"]:
            b64_str = row["photo_b64"]
            if "," in b64_str: b64_str = b64_str.split(",")[1]
            return send_file(io.BytesIO(base64.b64decode(b64_str)), mimetype="image/jpeg")
    except Exception: pass
    fallback = os.path.join(BASE_DIR, "static", "images", "icon-192.png")
    if os.path.exists(fallback): return send_file(fallback, mimetype="image/png")
    return jsonify({"error": "Not found"}), 404

@app.route("/student/<int:sid>/view")
@login_required
def view_student(sid):
    db = get_db()
    student = db.execute("SELECT * FROM students WHERE id=? AND is_active=1", (sid,)).fetchone()
    if not student: return redirect(url_for("students_page"))
    return render_template("view_student.html", student=student)

@app.route("/student/<int:sid>/edit", methods=["GET", "POST"])
@login_required
def edit_student(sid):
    db = get_db()
    student = db.execute("SELECT * FROM students WHERE id=? AND is_active=1", (sid,)).fetchone()
    if not student: return redirect(url_for("students_page"))
    if request.method == "POST":
        data = request.get_json() if request.is_json else request.form
        full_name, father_name = (data.get("full_name") or "").strip(), (data.get("father_name") or "").strip()
        class_name, section = (data.get("class_name") or "").strip(), (data.get("section") or "").strip()
        gr_number, phone = (data.get("gr_number") or "").strip(), (data.get("phone") or "").strip()
        if not full_name or not father_name or not class_name or not gr_number: return jsonify({"success": False, "error": "Fields missing"}), 400
        try:
            db.execute("UPDATE students SET full_name=?, father_name=?, class_name=?, section=?, gr_number=?, phone=? WHERE id=?", (full_name, father_name, class_name, section, gr_number, phone, sid))
            db.commit()
            log_audit("EDIT", "student", sid, f"Updated student GR#{gr_number} - {full_name}", session["user"], request.remote_addr)
            return jsonify({"success": True})
        except sqlite3.IntegrityError: return jsonify({"success": False, "error": "GR Number already exists."}), 400
    return render_template("edit_student.html", student=student, classes=Config.CLASSES, sections=Config.SECTIONS)

@app.route("/api/student/<int:sid>", methods=["DELETE"])
@login_required
def delete_student(sid):
    db = get_db()
    s = db.execute("SELECT full_name, gr_number FROM students WHERE id=?", (sid,)).fetchone()
    if s:
        db.execute("UPDATE students SET is_active=0 WHERE id=?", (sid,)); db.commit()
        log_audit("DELETE", "student", sid, f"Deleted GR#{s['gr_number']} - {s['full_name']}", session["user"], request.remote_addr)
    return jsonify({"success": True})


@app.route("/teachers")
@login_required
def teachers_page():
    db = get_db()
    teachers = db.execute("SELECT * FROM teachers WHERE is_active=1 ORDER BY full_name").fetchall()
    return render_template("teachers.html", teachers=teachers, count=len(teachers))

@app.route("/teacher/<int:tid>/photo")
@login_required
def teacher_photo(tid):
    db = get_db()
    file_path = os.path.join(BASE_DIR, "static", "uploads", "teachers", f"{tid}.jpg")
    if os.path.exists(file_path): return send_file(file_path, mimetype="image/jpeg")
    try:
        row = db.execute("SELECT photo_b64 FROM teachers WHERE id=?", (tid,)).fetchone()
        if row and row["photo_b64"]:
            b64_str = row["photo_b64"]
            if "," in b64_str: b64_str = b64_str.split(",")[1]
            return send_file(io.BytesIO(base64.b64decode(b64_str)), mimetype="image/jpeg")
    except Exception: pass
    fallback = os.path.join(BASE_DIR, "static", "images", "icon-192.png")
    if os.path.exists(fallback): return send_file(fallback, mimetype="image/png")
    return jsonify({"error": "Not found"}), 404

@app.route("/teacher/<int:tid>/view")
@login_required
def view_teacher(tid):
    db = get_db()
    teacher = db.execute("SELECT * FROM teachers WHERE id=? AND is_active=1", (tid,)).fetchone()
    if not teacher: return redirect(url_for("teachers_page"))
    return render_template("view_teacher.html", teacher=teacher)

@app.route("/teacher/<int:tid>/edit", methods=["GET", "POST"])
@login_required
def edit_teacher(tid):
    db = get_db()
    teacher = db.execute("SELECT * FROM teachers WHERE id=? AND is_active=1", (tid,)).fetchone()
    if not teacher: return redirect(url_for("teachers_page"))
    if request.method == "POST":
        data = request.get_json() if request.is_json else request.form
        full_name, father_name = (data.get("full_name") or "").strip(), (data.get("father_name") or "").strip()
        employee_id, designation = (data.get("employee_id") or "").strip(), (data.get("designation") or "").strip()
        phone = (data.get("phone") or "").strip()
        if not full_name or not employee_id: return jsonify({"success": False, "error": "Fields missing"}), 400
        try:
            db.execute("UPDATE teachers SET full_name=?, father_name=?, employee_id=?, designation=?, phone=? WHERE id=?", (full_name, father_name, employee_id, designation, phone, tid))
            db.commit()
            log_audit("EDIT", "teacher", tid, f"Updated EMP#{employee_id} - {full_name}", session["user"], request.remote_addr)
            return jsonify({"success": True})
        except sqlite3.IntegrityError: return jsonify({"success": False, "error": "EMP ID already exists."}), 400
    return render_template("edit_teacher.html", teacher=teacher)

@app.route("/api/teacher/<int:tid>", methods=["DELETE"])
@login_required
def delete_teacher(tid):
    db = get_db()
    t = db.execute("SELECT full_name, employee_id FROM teachers WHERE id=?", (tid,)).fetchone()
    if t:
        db.execute("UPDATE teachers SET is_active=0 WHERE id=?", (tid,)); db.commit()
        log_audit("DELETE", "teacher", tid, f"Deleted EMP#{t['employee_id']} - {t['full_name']}", session["user"], request.remote_addr)
    return jsonify({"success": True})


# ==========================================
# SCANNERS (ZERO HANG - LIVENESS ENABLED)
# ==========================================
@app.route("/scanner-students")
@login_required
def scanner_students():
    return render_template("scanner_students.html")

@app.route("/api/scan-student", methods=["POST"])
@login_required
def scan_student():
    data = request.get_json() or {}
    frame_b64 = data.get("frame", "")
    if not frame_b64: return jsonify({"status": "unknown"})
    
    db = get_db()
    pkt_now = get_pkt_now()
    today = pkt_now.date().isoformat()
    now = pkt_now.strftime("%H:%M:%S")
    
    if db.execute("SELECT reason FROM holidays WHERE date=?", (today,)).fetchone() or pkt_now.weekday() == 6:
        return jsonify({"status": "holiday", "message": "Today is a holiday."})
        
    known_list = [(r["id"], "student", r["face_embedding"]) for r in db.execute("SELECT id, face_embedding FROM students WHERE is_active=1 AND face_embedding IS NOT NULL").fetchall()]
    if not known_list: return jsonify({"status": "unknown"})
    
    sid, ftype, face_box = recognize_face(frame_b64, known_list)
    if sid is None: return jsonify({"status": "unknown"})
    student = db.execute("SELECT * FROM students WHERE id=?", (sid,)).fetchone()
    
    try:
        frame_img = decode_base64_image(frame_b64)
        strictness = get_setting("liveness_strictness", "High")
        liveness = check_liveness(frame_img, face_box, strictness)
        if not liveness["live"]:
            reason_str = ", ".join(liveness["reasons"])
            log_audit("SPOOF_BLOCKED", "security", sid, f"Fake face blocked for {student['full_name']}: {reason_str}", session.get("user", "system"), request.remote_addr)
            return jsonify({"status": "spoof", "message": "Fake face blocked!"})
    except Exception: pass
    
    existing = db.execute("SELECT scan_time FROM student_attendance WHERE student_id=? AND date=?", (sid, today)).fetchone()
    if existing: return jsonify({"status": "duplicate", "name": student["full_name"], "time": existing["scan_time"]})
    
    cutoff = get_setting("student_late_cutoff", Config.STUDENT_LATE_CUTOFF)
    status = "PRESENT" if now <= cutoff else "LATE"
    db.execute("INSERT INTO student_attendance (student_id, date, scan_time, status) VALUES (?,?,?,?)", (sid, today, now, status))
    db.commit()
    return jsonify({"status": "success", "name": student["full_name"], "class": student["class_name"], "time": now, "attendance_status": status})


@app.route("/scanner-teachers")
@login_required
def scanner_teachers():
    return render_template("scanner_teachers.html")

@app.route("/api/scan-teacher", methods=["POST"])
@login_required
def scan_teacher():
    data = request.get_json() or {}
    frame_b64 = data.get("frame", "")
    if not frame_b64: return jsonify({"status": "unknown"})
    
    db = get_db()
    pkt_now = get_pkt_now()
    today = pkt_now.date().isoformat()
    now = pkt_now.strftime("%H:%M:%S")
    
    if db.execute("SELECT reason FROM holidays WHERE date=?", (today,)).fetchone() or pkt_now.weekday() == 6:
        return jsonify({"status": "holiday", "message": "Today is a holiday."})
        
    known_list = [(r["id"], "teacher", r["face_embedding"]) for r in db.execute("SELECT id, face_embedding FROM teachers WHERE is_active=1 AND face_embedding IS NOT NULL").fetchall()]
    if not known_list: return jsonify({"status": "unknown"})
    
    tid, ftype, face_box = recognize_face(frame_b64, known_list)
    if tid is None: return jsonify({"status": "unknown"})
    teacher = db.execute("SELECT * FROM teachers WHERE id=?", (tid,)).fetchone()
    
    existing = db.execute("SELECT * FROM teacher_attendance WHERE teacher_id=? AND date=?", (tid, today)).fetchone()
    if existing and existing["check_out"] is not None:
        return jsonify({"status": "already_out", "name": teacher["full_name"], "message": f"Already checked out at {existing['check_out']}"})
        
    try:
        frame_img = decode_base64_image(frame_b64)
        strictness = get_setting("liveness_strictness", "High")
        liveness = check_liveness(frame_img, face_box, strictness)
        if not liveness["live"]:
            reason_str = ", ".join(liveness["reasons"])
            log_audit("SPOOF_BLOCKED", "security", tid, f"Fake face blocked for {teacher['full_name']}: {reason_str}", session.get("user", "system"), request.remote_addr)
            return jsonify({"status": "spoof", "message": "Fake face blocked!"})
    except Exception: pass
    
    if existing is None:
        status = "PRESENT" if now <= get_setting("teacher_late_cutoff", Config.TEACHER_LATE_CUTOFF) else "LATE"
        db.execute("INSERT INTO teacher_attendance (teacher_id, date, check_in, status) VALUES (?,?,?,?)", (tid, today, now, status))
        db.commit()
        return jsonify({"status": "checkin", "name": teacher["full_name"], "time": now, "attendance_status": status})
        
    dup_interval = int(get_setting("duplicate_scan_interval", "30"))
    diff = datetime.strptime(now, "%H:%M:%S") - datetime.strptime(existing["check_in"], "%H:%M:%S")
    if diff.total_seconds() < dup_interval:
        return jsonify({"status": "duplicate", "name": teacher["full_name"], "message": f"Wait {dup_interval}s to check out."})
        
    duty = f"{int(diff.total_seconds()//3600)}h {int((diff.total_seconds()%3600)//60)}m"
    status = "HALF-DAY" if now < get_setting("teacher_half_day_cutoff", Config.TEACHER_HALF_DAY_CUTOFF) else "FULL-DAY" if now >= get_setting("teacher_full_day_cutoff", Config.TEACHER_FULL_DAY_CUTOFF) else existing["status"]
    
    db.execute("UPDATE teacher_attendance SET check_out=?, duty_hours=?, status=? WHERE id=?", (now, duty, status, existing["id"]))
    db.commit()
    return jsonify({"status": "checkout", "name": teacher["full_name"], "time": now, "duty_hours": duty})


# ==========================================
# HISTORY, REPORTS, SETTINGS & GUIDE
# ==========================================
@app.route("/install")
@login_required
def install_page():
    return render_template("install.html")

@app.route("/guide")
@login_required
def guide_page():
    return render_template("guide.html")

@app.route("/api/save-absence-reason", methods=["POST"])
@login_required
def save_absence_reason():
    data = request.get_json() or {}
    tid, tname, reason = data.get("teacher_id"), data.get("teacher_name"), data.get("reason")
    pkt_now = get_pkt_now()
    today, day_name = pkt_now.date().isoformat(), pkt_now.strftime("%A")
    db = get_db()
    if not db.execute("SELECT id FROM teacher_absences WHERE teacher_id=? AND date=?", (tid, today)).fetchone():
        db.execute("INSERT INTO teacher_absences (teacher_id, teacher_name, date, day_name, reason) VALUES (?,?,?,?,?)", (tid, tname, today, day_name, reason))
        db.commit()
    return jsonify({"success": True})

@app.route("/absences")
@login_required
def absences_page():
    db = get_db()
    records = db.execute("SELECT * FROM teacher_absences ORDER BY timestamp DESC").fetchall()
    return render_template("absences.html", records=records)

@app.route("/api/download-absences")
@login_required
def download_absences():
    db = get_db()
    records = db.execute("SELECT teacher_name, date, day_name, reason, timestamp FROM teacher_absences ORDER BY timestamp DESC").fetchall()
    wb = Workbook()
    ws = wb.active
    ws.title = "Absence Records"
    headers = ["Teacher Name", "Date", "Day", "Reason", "Submitted At"]
    for col, h in enumerate(headers, 1): ws.cell(row=1, column=col, value=h).font = Font(bold=True)
    for i, r in enumerate(records, 2):
        ws.cell(row=i, column=1, value=r["teacher_name"]); ws.cell(row=i, column=2, value=r["date"]); ws.cell(row=i, column=3, value=r["day_name"]); ws.cell(row=i, column=4, value=r["reason"]); ws.cell(row=i, column=5, value=r["timestamp"])
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", as_attachment=True, download_name="Absent_Teachers_Record.xlsx")

@app.route("/attendance-history")
@login_required
def attendance_history():
    db = get_db()
    selected_date = request.args.get("date", get_pkt_now().date().isoformat())
    student_att = db.execute("SELECT s.full_name, s.gr_number, s.class_name, s.section, sa.scan_time, sa.status FROM student_attendance sa JOIN students s ON s.id = sa.student_id WHERE sa.date = ? ORDER BY s.class_name, sa.scan_time", (selected_date,)).fetchall()
    teacher_att = db.execute("SELECT t.full_name, t.employee_id, t.designation, ta.check_in, ta.check_out, ta.duty_hours, ta.status FROM teacher_attendance ta JOIN teachers t ON t.id = ta.teacher_id WHERE ta.date = ? ORDER BY ta.check_in", (selected_date,)).fetchall()
    return render_template("attendance_history.html", student_att=student_att, teacher_att=teacher_att, selected_date=selected_date)

@app.route("/history")
@login_required
def history_page():
    db = get_db()
    filter_type = request.args.get("type", "REGISTER").upper()
    if filter_type not in ("REGISTER", "DELETE"): filter_type = "REGISTER"
    rows = db.execute("SELECT * FROM audit_log WHERE action=? ORDER BY timestamp DESC LIMIT 500", (filter_type,)).fetchall()
    logs = []
    for r in rows:
        dt = datetime.fromisoformat(r["timestamp"]) if "T" in r["timestamp"] else datetime.strptime(r["timestamp"], "%Y-%m-%d %H:%M:%S")
        logs.append({"date_only": dt.strftime("%Y-%m-%d"), "day_name": dt.strftime("%A"), "time_only": dt.strftime("%H:%M:%S"), "target_type": r["target_type"] or "-", "details": r["details"], "admin_user": r["admin_user"]})
    return render_template("history.html", logs=logs, filter_type=filter_type)

@app.route("/reports")
@login_required
def reports_page():
    return render_template("reports.html")

@app.route("/api/report-daily-excel")
@login_required
def report_daily_excel():
    db = get_db()
    today = get_pkt_now().date().isoformat()
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Student Attendance"
    for col, h in enumerate(["GR#", "Name", "Father", "Class", "Section", "Time", "Status"], 1): ws1.cell(row=1, column=col, value=h)
    for i, r in enumerate(db.execute("SELECT s.gr_number, s.full_name, s.father_name, s.class_name, s.section, sa.scan_time, sa.status FROM student_attendance sa JOIN students s ON s.id = sa.student_id WHERE sa.date = ?", (today,)).fetchall(), 2):
        for j, val in enumerate([r["gr_number"], r["full_name"], r["father_name"], r["class_name"], r["section"], r["scan_time"], r["status"]], 1): ws1.cell(row=i, column=j, value=val)
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", as_attachment=True, download_name=f"BSS_{today}.xlsx")

@app.route("/api/report-monthly-pdf")
@login_required
def report_monthly_pdf():
    db = get_db()
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=A4)
    doc.build([Paragraph("M.ALI Schools - Monthly Report", getSampleStyleSheet()["Title"]), Spacer(1, 20)])
    output.seek(0)
    return send_file(output, mimetype="application/pdf", as_attachment=True, download_name="BSS_Monthly.pdf")

@app.route("/api/admin/clear-data", methods=["POST"])
@login_required
def clear_admin_data():
    data = request.get_json() or {}
    target = data.get("target")
    password = data.get("password", "")
    if password != CLEAR_DATA_PASSWORD:
        return jsonify({"success": False, "message": "Incorrect Admin Password!"}), 401

    db = get_db()
    today = get_pkt_now().date().isoformat()
    current_username = session.get("user", "admin")

    try:
        if target == "dashboard":
            db.execute("DELETE FROM student_attendance WHERE date=?", (today,))
            db.execute("DELETE FROM teacher_attendance WHERE date=?", (today,))
            message = "Dashboard daily scans cleared!"
        elif target == "history":
            db.execute("DELETE FROM audit_log WHERE action IN ('REGISTER', 'DELETE')")
            message = "History cleared!"
        elif target == "settings":
            db.execute("DELETE FROM audit_log")
            message = "Settings Audit Logs cleared!"
        else: return jsonify({"success": False, "message": "Invalid target."}), 400
        db.commit()
        log_audit("CLEAR_DATA", "system", None, f"Admin cleared data for: {target}", current_username, request.remote_addr)
        return jsonify({"success": True, "message": message})
    except Exception as e:
        db.execute("ROLLBACK")
        return jsonify({"success": False, "message": f"Database error: {str(e)}"}), 500

@app.route("/settings", methods=["GET", "POST"])
@login_required
def settings_page():
    db = get_db()
    if request.method == "POST":
        data = request.get_json() or {}
        action = data.get("action", "")
        if action == "change_password":
            user = db.execute("SELECT * FROM users WHERE username=?", (session["user"],)).fetchone()
            if user and bcrypt.checkpw(data.get("old_password", "").encode(), user["password_hash"].encode()):
                db.execute("UPDATE users SET password_hash=? WHERE id=?", (bcrypt.hashpw(data.get("new_password", "").encode(), bcrypt.gensalt()).decode(), user["id"]))
                db.commit()
                return jsonify({"success": True})
            return jsonify({"success": False, "error": "Old password incorrect."}), 400
        if action == "update_setting":
            db.execute("INSERT OR REPLACE INTO settings (key, value) VALUES (?,?)", (data.get("key"), data.get("value")))
            db.commit()
            return jsonify({"success": True})
        if action == "add_holiday":
            db.execute("INSERT INTO holidays (date, reason, created_by) VALUES (?,?,?)", (data.get("date"), data.get("reason"), session["user"]))
            db.commit(); return jsonify({"success": True})
        if action == "delete_holiday":
            db.execute("DELETE FROM holidays WHERE id=?", (data.get("id"),))
            db.commit(); return jsonify({"success": True})
        if action == "remove_device":
            db.execute("DELETE FROM authorized_devices WHERE id=?", (data.get("id"),))
            db.commit(); return jsonify({"success": True})

    return render_template("settings.html", 
        settings={r["key"]: r["value"] for r in db.execute("SELECT * FROM settings").fetchall()},
        devices=db.execute("SELECT * FROM authorized_devices ORDER BY authorized_at DESC").fetchall(),
        holidays=db.execute("SELECT * FROM holidays ORDER BY date DESC").fetchall(),
        audit=db.execute("SELECT * FROM audit_log ORDER BY timestamp DESC LIMIT 20").fetchall())

@app.route("/api/upload-logo", methods=["POST"])
@login_required
def upload_logo():
    try:
        f = request.files.get("logo")
        if not f: return jsonify({"success": False, "error": "No file uploaded."}), 400
        img = Image.open(f.stream).convert("RGBA")
        out_dir = os.path.join(BASE_DIR, "static", "images")
        os.makedirs(out_dir, exist_ok=True)
        img.resize((192, 192), Image.LANCZOS).save(os.path.join(out_dir, "icon-192.png"), "PNG")
        img.resize((512, 512), Image.LANCZOS).save(os.path.join(out_dir, "icon-512.png"), "PNG")
        log_audit("UPLOAD_LOGO", "system", None, "Custom logo uploaded", session["user"], request.remote_addr)
        return jsonify({"success": True})
    except Exception as e: return jsonify({"success": False, "error": str(e)}), 500

@app.route("/manifest.json")
def manifest(): return send_file(os.path.join(BASE_DIR, "static", "manifest.json"), mimetype="application/json")
@app.route("/sw.js")
def service_worker(): return send_file(os.path.join(BASE_DIR, "static", "sw.js"), mimetype="application/javascript")

if __name__ == "__main__":
    init_db()
    app.run(host="0.0.0.0", port=5000, debug=False, threaded=True)